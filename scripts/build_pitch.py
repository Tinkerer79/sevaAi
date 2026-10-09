"""Build the SW-01 pitch deck from the supplied hackathon template."""
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / 'Presentation Template.pptx'
OUT = ROOT / 'SevaManipur_AI4SEVA_SW01.pptx'
SHOT = ROOT / 'gui-test-screenshots' / 't1_home_hero.png'

NAVY = '10234B'
BLUE = '1769E0'
CYAN = '22B8D6'
MINT = '7CE0C3'
INK = '182742'
MUTED = '5D6C82'
PALE = 'F3F7FC'
LINE = 'DCE5F0'
WHITE = 'FFFFFF'
GOLD = 'FFC45B'

prs = Presentation(str(TEMPLATE))
W, H = prs.slide_width / 914400, prs.slide_height / 914400


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def remove_content(slide):
    tree = slide.shapes._spTree
    for shape in list(slide.shapes):
        tree.remove(shape._element)


def rect(slide, x, y, w, h, fill, line=None, radius=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line or fill)
    if radius:
        shape.adjustments[0] = 0.12
    return shape


def text(slide, value, x, y, w, h, size=16, color=INK, bold=False,
         font='Aptos', align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
         margin=0, italic=False, fit=True):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    if fit:
        from pptx.enum.text import MSO_AUTO_SIZE
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    return box


def label(slide, value, x, y, w, color=BLUE):
    text(slide, value.upper(), x, y, w, 0.22, 9, color, True)


def card(slide, x, y, w, h, title, body, accent=BLUE, title_size=15, body_size=12.5):
    rect(slide, x, y, w, h, WHITE, LINE, True)
    rect(slide, x, y, 0.06, h, accent, accent, True)
    text(slide, title, x + 0.22, y + 0.19, w - 0.4, 0.36, title_size, INK, True)
    text(slide, body, x + 0.22, y + 0.66, w - 0.42, h - 0.78, body_size, MUTED)


def footer(slide, n, dark=False):
    col = 'B8C9E3' if dark else '8795A8'
    text(slide, 'SEVAMANIPUR AI  ·  AI4SEVA 2026  ·  PROTOTYPE', 0.62, H - 0.34, 7.0, 0.17, 8, col, True)
    text(slide, f'{n:02d} / 06', W - 1.25, H - 0.34, 0.65, 0.17, 8, col, True, align=PP_ALIGN.RIGHT)


def header(slide, n, title_value, kicker):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb('F8FAFD')
    rect(slide, 0, 0, W, 0.12, BLUE)
    label(slide, f'{n:02d}  /  {kicker}', 0.66, 0.38, 7)
    text(slide, title_value, 0.66, 0.72, W - 1.32, 0.64, 27, NAVY, True)
    footer(slide, n)


for s in prs.slides:
    remove_content(s)

# 1 — Problem statement / cover
s = prs.slides[0]
s.background.fill.solid()
s.background.fill.fore_color.rgb = rgb(NAVY)
rect(s, 0, 0, 0.15, H, CYAN)
label(s, 'AI4SEVA 2026  ·  SOCIAL WELFARE  ·  SW-01', 0.72, 0.58, 6.7, MINT)
text(s, 'Find support.\nKnow your next step.', 0.72, 1.16, 6.3, 1.62, 34, WHITE, True)
text(s, 'SevaManipur AI', 0.75, 3.02, 5.8, 0.55, 24, MINT, True)
text(s, 'SW-01 · Welfare Scheme Navigator', 0.75, 3.62, 5.9, 0.45, 17, 'E0EAF8', False)
text(s, 'Find possible matches, understand requirements and plan what to do next.', 0.75, 4.28, 5.6, 0.72, 13, 'B8C9E3')
rect(s, 0.75, 5.75, 2.88, 0.76, '1A315D', '49658E', True)
text(s, 'PITCH PRESENTATION', 0.96, 5.91, 2.5, 0.2, 9, 'D7E5F7', True)
text(s, 'SW-01 · SOCIAL WELFARE', 0.96, 6.16, 2.5, 0.2, 10, WHITE, True)
text(s, '9 OCTOBER 2026', 0.76, 6.76, 2.5, 0.2, 8, 'AFC1DD', True)
rect(s, 7.55, 1.02, 5.1, 3.17, WHITE, '536B91', True)
s.shapes.add_picture(str(SHOT), Inches(7.62), Inches(1.09), width=Inches(4.96), height=Inches(2.79))
text(s, 'Citizen gateway web app prototype', 7.72, 3.93, 4.75, 0.2, 9, MUTED, True)
rect(s, 7.76, 4.65, 4.2, 0.72, '1A315D', '385982', True)
text(s, 'PROFILE  →  POSSIBLE MATCHES  →  NEXT STEPS', 7.85, 4.89, 4.0, 0.2, 10, MINT, True, align=PP_ALIGN.CENTER)
text(s, 'Prototype developed for the AI4SEVA Hackathon. Not an official Government of Manipur website.', 7.72, 5.66, 4.65, 0.56, 10, 'BACAE2')
footer(s, 1, True)

# 2 — Problem and proposed solution
s = prs.slides[1]
header(s, 2, 'Help citizens move from questions to next steps', 'PROBLEM & SOLUTION')
text(s, 'THE CITIZEN PROBLEM', 0.72, 1.65, 4, 0.24, 10, MUTED, True)
card(s, 0.7, 2.0, 3.72, 1.52, '“What might fit me?”', 'Scheme rules can be hard to find and compare.', BLUE)
card(s, 4.56, 2.0, 3.72, 1.52, '“What do I need?”', 'Documents and application steps are spread across sources.', CYAN)
card(s, 8.42, 2.0, 3.9, 1.52, '“What happens next?”', 'People need guidance, while the department makes the final decision.', GOLD)
text(s, 'OUR RESPONSE', 0.72, 3.92, 3, 0.24, 10, MUTED, True)
steps = [
    ('01', 'Share basics', 'Age · income · work · area'),
    ('02', 'See possible matches', 'Basic rule checks + reasons'),
    ('03', 'Plan next steps', 'Documents · process · official links'),
]
for i, (num, title_value, body) in enumerate(steps):
    x = 0.72 + i * 4.02
    rect(s, x, 4.32, 3.64, 1.25, WHITE, LINE, True)
    rect(s, x + 0.18, 4.55, 0.48, 0.48, BLUE if i != 1 else CYAN, None, True)
    text(s, num, x + 0.18, 4.68, 0.48, 0.16, 10, WHITE, True, align=PP_ALIGN.CENTER)
    text(s, title_value, x + 0.82, 4.51, 2.6, 0.28, 14, INK, True)
    text(s, body, x + 0.82, 4.91, 2.55, 0.36, 10.5, MUTED)
    if i < 2:
        text(s, '→', x + 3.7, 4.73, 0.28, 0.3, 18, BLUE, True, align=PP_ALIGN.CENTER)
rect(s, 0.72, 5.92, 11.5, 0.6, 'EAF5F7', 'D4EBEF', True)
text(s, 'The app offers guidance. The department decides eligibility.', 0.96, 6.1, 11.05, 0.22, 12, NAVY, True, align=PP_ALIGN.CENTER)

# 3 — Differentiation and safeguards
s = prs.slides[2]
header(s, 3, 'Clear rules. Useful explanations. Human decisions.', 'HOW IT WORKS')
text(s, 'PROFILE', 0.8, 1.82, 1.5, 0.2, 9, BLUE, True)
text(s, 'RULE MATCH', 3.62, 1.82, 1.8, 0.2, 9, CYAN, True)
text(s, 'WHY THIS MATCH', 6.58, 1.82, 2.1, 0.2, 9, BLUE, True)
text(s, 'NEXT ACTION', 9.77, 1.82, 1.7, 0.2, 9, CYAN, True)
for x, title_value, body, accent in [
    (0.72, 'Citizen inputs', 'Age, occupation,\nincome and area', BLUE),
    (3.48, 'Deterministic checks', 'Transparent scheme\ncriteria filters', CYAN),
    (6.4, 'Explainable result', 'Show match reasons,\nsource and caveats', BLUE),
    (9.55, 'Guided application', 'Documents, process\nand AI assistance', CYAN),
]:
    rect(s, x, 2.18, 2.42, 1.25, WHITE, LINE, True)
    rect(s, x, 2.18, 2.42, 0.08, accent, accent, True)
    text(s, title_value, x + 0.15, 2.45, 2.12, 0.3, 13, INK, True, align=PP_ALIGN.CENTER)
    text(s, body, x + 0.15, 2.84, 2.12, 0.48, 11, MUTED, align=PP_ALIGN.CENTER)
for x in [3.13, 6.06, 9.12]:
    text(s, '→', x, 2.59, 0.28, 0.32, 18, BLUE, True, align=PP_ALIGN.CENTER)
card(s, 0.72, 3.92, 3.72, 1.55, 'Basic rule checks', 'The current matcher checks age, work, area and income.', BLUE)
card(s, 4.56, 3.92, 3.72, 1.55, 'AI explains the catalog', 'The assistant explains available records and points to listed sources.', CYAN)
card(s, 8.42, 3.92, 3.9, 1.55, 'Department decides', 'A match is not approval. Confirm current rules and application status with the department.', GOLD)
text(s, 'The matcher is a prototype: scheme-specific checks, such as disability or widow status, still need to be added.', 0.8, 5.92, 11.45, 0.55, 10.5, MUTED, italic=True)

# 4 — Product/model overview
s = prs.slides[3]
header(s, 4, 'What works today — and what needs validation', 'THE PROTOTYPE')
rect(s, 0.72, 1.68, 5.7, 3.2, WHITE, LINE, True)
s.shapes.add_picture(str(SHOT), Inches(0.82), Inches(1.78), width=Inches(5.5), height=Inches(3.09))
text(s, 'Current home screen · live typed question handoff to Seva AI', 0.92, 4.99, 5.2, 0.22, 9, MUTED, True)
label(s, 'CURRENT PROTOTYPE', 6.85, 1.72, 4, BLUE)
card(s, 6.82, 2.08, 2.57, 1.28, 'Scheme finder', 'Basic checks: age, work, area and income', BLUE, 13, 10.5)
card(s, 9.56, 2.08, 2.74, 1.28, 'Seva AI', 'Explains catalog information in a chat', CYAN, 13, 10.5)
card(s, 6.82, 3.54, 2.57, 1.28, 'Languages', 'English, Hindi and Meiteilon controls; translation review remains', CYAN, 12, 10)
card(s, 9.56, 3.54, 2.74, 1.28, 'Live architecture', 'Vercel app and APIs; Supabase catalog and signed-in data', BLUE, 12, 10)
rect(s, 6.84, 5.08, 5.44, 0.84, 'FFF6E6', 'F4D99F', True)
text(s, 'DATA STATUS', 7.02, 5.23, 1.25, 0.2, 9, '885D12', True)
text(s, '7 source-linked Social Welfare records are in Supabase; the department must confirm current rules and availability. Other records are samples.', 8.2, 5.15, 3.82, 0.58, 9.6, '704D18')
text(s, 'Next: validate the catalog, add scheme-specific rules and test with citizens.', 0.78, 6.28, 11.4, 0.32, 11, NAVY, True)

# 5 — Sustainability
s = prs.slides[4]
header(s, 5, 'A practical pilot needs a trusted owner', 'PILOT & SUSTAINABILITY')
text(s, 'PUBLIC VALUE FIRST', 0.76, 1.66, 4, 0.24, 10, BLUE, True)
text(s, 'Keep scheme discovery and guidance free for citizens.', 0.76, 2.03, 5.0, 0.72, 21, NAVY, True)
text(s, 'SUSTAINABILITY MODEL', 6.48, 1.66, 4, 0.24, 10, CYAN, True)
card(s, 6.44, 2.03, 2.75, 1.44, 'Content owner', 'Department reviews rules, sources and links.', BLUE, 13, 10.5)
card(s, 9.36, 2.03, 2.86, 1.44, 'Review schedule', 'Check changes and application status regularly.', CYAN, 13, 10.5)
card(s, 0.76, 3.13, 2.75, 1.44, 'Pilot & learn', 'Test whether citizens reach the right next step.', CYAN, 13, 10.5)
card(s, 3.68, 3.13, 2.75, 1.44, 'Protect users', 'Minimize personal data and review privacy needs.', BLUE, 13, 10.5)
card(s, 6.44, 3.82, 2.75, 1.44, 'Fund operations', 'Hosting and maintenance need an approved owner and budget.', BLUE, 13, 10.3)
card(s, 9.36, 3.82, 2.86, 1.44, 'Expand in stages', 'Add departments after their records are reviewed.', CYAN, 13, 10.5)
rect(s, 0.78, 5.4, 11.42, 0.7, NAVY, NAVY, True)
text(s, 'A real service needs an accountable content owner and a regular review cycle.', 1.0, 5.62, 10.95, 0.24, 12, WHITE, True, align=PP_ALIGN.CENTER)
text(s, 'This is a proposed pilot model; no government partnership or funding is claimed.', 0.82, 6.28, 11, 0.25, 9.5, MUTED, italic=True)

# 6 — Market and scalability
s = prs.slides[5]
header(s, 6, 'Start with Social Welfare. Expand after review.', 'ROLLOUT PLAN')
text(s, 'A focused entry point', 0.75, 1.66, 4.3, 0.4, 19, NAVY, True)
text(s, 'Keep state and central schemes clear. Add departments only after checking their information.', 0.75, 2.12, 4.6, 0.67, 13, MUTED)
for i, (title_value, sub, color) in enumerate([
    ('ALL', 'Browse full catalog', BLUE),
    ('MANIPUR', 'State schemes', CYAN),
    ('OTHER SCHEMES', 'Central / other', NAVY),
]):
    x = 0.78 + i * 1.62
    rect(s, x, 3.02, 1.48, 0.88, color, color, True)
    text(s, title_value, x + 0.06, 3.2, 1.36, 0.2, 9, WHITE, True, align=PP_ALIGN.CENTER)
    text(s, sub, x + 0.08, 3.5, 1.32, 0.18, 8, 'E6F4FF', align=PP_ALIGN.CENTER)
text(s, '+ Filter by department', 0.8, 4.15, 4.5, 0.3, 12, BLUE, True)
label(s, 'PILOT IN PHASES', 6.1, 1.68, 3, CYAN)
card(s, 6.06, 2.04, 2.0, 1.23, '01  Verify', 'Social Welfare sources + rules', BLUE, 12, 10)
card(s, 8.18, 2.04, 2.0, 1.23, '02  Pilot', 'Test with citizens and staff', CYAN, 12, 10)
card(s, 10.3, 2.04, 2.0, 1.23, '03  Expand', 'Add reviewed departments', NAVY, 12, 10)
text(s, 'Built for careful growth', 6.1, 3.72, 4, 0.3, 14, NAVY, True)
text(s, '• One catalog with sources and review dates\n• Simple matching rules, separate from AI explanations\n• Vercel app with Supabase catalog and user data', 6.12, 4.16, 6.08, 1.38, 12, MUTED)
rect(s, 0.78, 5.18, 4.76, 0.93, 'EAF5F7', 'D4EBEF', True)
text(s, 'FIRST DATA DROP', 1.0, 5.36, 1.5, 0.18, 9, BLUE, True)
text(s, '7 Social Welfare records link to official sources; current rules and application status need department confirmation.', 1.0, 5.61, 4.27, 0.37, 9.7, NAVY)
text(s, 'SevaManipur AI is a prototype developed for AI4SEVA 2026; not an official government service.', 6.12, 6.0, 6.05, 0.38, 10, NAVY, True)

prs.core_properties.title = 'SevaManipur AI — SW-01'
prs.core_properties.subject = 'AI-based Welfare Scheme Eligibility & Assistance Navigator'
prs.core_properties.author = 'SevaManipur AI Hackathon Team'
prs.core_properties.keywords = 'AI4SEVA, Manipur, Social Welfare, SW-01'
prs.save(str(OUT))
print(OUT)
